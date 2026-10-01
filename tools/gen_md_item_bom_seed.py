"""
gen_md_item_bom_seed.py — docs/260828 BOM Master List.xlsx → dist/seed_md_item_bom_master_list.sql

엑셀(공장별 시트 · ■ 차종 헤더 · Complete ASSY 블록 · L1~L3 들여쓰기 BOM)을 읽어
MD_Item / MD_BomVersion / MD_Bom / MD_Vendor / SCM_ItemVendor 전체 재적재 SQL 을 만든다.
엑셀이 유일한 정본이다 — 시드는 기존 MD_Item·MD_Bom 을 모두 지우고 넣는다.

사용:  python tools/gen_md_item_bom_seed.py [엑셀경로] [출력경로]
"""
from __future__ import annotations

import collections
import hashlib
import re
import sys
from pathlib import Path

import openpyxl
from openpyxl.utils import column_index_from_string as ci

ROOT = Path(__file__).resolve().parents[1]
XLSX = Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / "docs" / "260828 BOM Master List.xlsx"
OUT = Path(sys.argv[2]) if len(sys.argv) > 2 else ROOT / "dist" / "seed_md_item_bom_master_list.sql"

EFF_FROM = "2026-08-28"          # 엑셀 파일 날짜
CREATED_BY = "seed"
MAX_NAME = 80

# ── 원본 오류 보정 ─────────────────────────────────────────────────────────
# (시트, 행) → 올바른 루트 품번. 자식 품번(LH/RH)·이웃 블록과 맞춰 판단.
ROOT_FIX = {
    ("서베너", 31): "M3320-PI020",   # RR RH 블록이 LH 와 같은 M3310-PI020 로 적혀 있음
    ("서베너", 97): "M2320-QI000",   # FR RH 루트가 자기 L1 자식 82320-QI000 과 같음
    ("서베너", 135): "M2320-FC000",  # FR RH 블록 루트가 LH 변형 2 (M2310-FC010) 로 적혀 있음
}
COLOR_FIX = {"YUG": "YGU"}         # NEA PE 조립색 오타 (DB·SEMS 는 YGU)
NO_COLOR = {"N/C"}
SUPPLIER_FIX = {"KD.HANHWA": "KD.HANWHA"}

# 변형 2 이상 루트의 품명 (블록 제목은 변형 1 기준이라 자식 품명에서 따온다)
VARIANT_ROOT_NAME = {
    "M3311-TD100": "MODULE ASSY-RR DR TRIM UPR,LH(+CUR)",
    "M3321-TD100": "MODULE ASSY-RR DR TRIM UPR,RH(+CUR)",
    "M83371-P8010": "MODULE RR DR TRIM UPR, LH (FAKE STITCH, +CUR)",
    "M83381-P8010": "MODULE RR DR TRIM UPR, RH (FAKE STITCH, +CUR)",
    "M2310-FC010": "PNL ASSY-FR DR TRIM UPR, LH (FAKE STITCH, IMS_LHD)",
    "M3310-FC010": "PNL ASSY-RR DR TRIM UPR, LH (FAKE STITCH, CUR)",
    "M3320-FC010": "PNL ASSY-RR DR TRIM UPR, RH (FAKE STITCH, CUR)",
}

SECTION_CAR = {
    "NEA PE(서베너)": "NE1A", "ME1a(서베너)": "ME1A", "NE1A W(서베너)": "NE1A", "LX3a(서베너)": "LX3A",
    "LQ2(조지아)": "LQ2", "NQ5a(조지아)": "NQ5A", "MV1a(조지아)": "MV1A", "NX5A(어번)": "NX5A",
}

# 공급사(AU 열) → MD_Vendor. MIP.JACKSON 은 자사 공장이라 벤더가 아니다(SUB 판정에만 쓴다).
VENDORS = {
    "LP.DEAHA":               ("LP-DEAHA",     "DEAHA",               "LOCAL", "Resin"),
    "LP.KRA":                 ("LP-KRA",       "KRA",                 "LOCAL", "Interior Trim"),
    "LP.HYUNDAI_EP":          ("LP-HYUNDAIEP", "HYUNDAI EP",          "LOCAL", "Resin"),
    "KD.HANWHA":              ("KD-HANWHA",    "HANWHA",              "CKD",   "Fabric"),
    "KD.KOLON_GLOTECH":       ("KD-KOLON",     "KOLON GLOTECH",       "CKD",   "Fabric"),
    "KD.SIKA":                ("KD-SIKA",      "SIKA",                "CKD",   "Adhesive"),
    "KD.KLUEBER":             ("KD-KLUEBER",   "KLUEBER",             "CKD",   "Lubricant"),
    "KD.SUNGCHANG_PRECISION": ("KD-SUNGCHANG", "SUNGCHANG PRECISION", "CKD",   "Fastener"),
}
IN_HOUSE = "MIP"

# 신설 ITEM_CATEGORY (Attribute1 = 허용 ItemType — MD-003 콤보 게이트)
NEW_CATEGORIES = [
    ("RESIN",    "레진",     "Resin",     "MATERIAL", 94),
    ("CHEM",     "부자재",   "Chemical",  "MATERIAL", 95),
    ("FASTENER", "체결부품", "Fastener",  "MATERIAL", 96),
    ("PART",     "구매부품", "Part",      "MATERIAL", 97),
]

# 2026-09-24 AMES_DEV 스냅샷 — 엑셀에 없는 PGN·ALC(IMG 완제품 라벨이 쓴다)를 되살린다.
PGN_ALC = {
    '82311-TD000NNB': ('QSUB', 'M094'), '82311-TD000PNY': ('QSUB', 'M096'),
    '82311-TD000VKE': ('QSUB', 'M097'), '82311-TD000YGN': ('QSUB', 'M095'),
    '83311-TD000NNB': ('QSUB', 'M102'), '83311-TD000PNY': ('QSUB', 'M105'),
    '83311-TD000VKE': ('QSUB', 'M107'), '83311-TD000YGN': ('QSUB', 'M103'),
    '83311-TD100NNB': ('QSUB', 'M104'), '83311-TD100PNY': ('QSUB', 'M106'),
    '83311-TD100VKE': ('QSUB', 'M108'), '83311-TD100YGN': ('QSUB', 'M109'),
    '83314-P8000': ('AQFG', 'DL001'), '83314-P8010': ('AQFG', 'DL003'),
    '83321-TD000NNB': ('QSUB', 'M110'), '83321-TD000PNY': ('QSUB', 'M113'),
    '83321-TD000VKE': ('QSUB', 'M115'), '83321-TD000YGN': ('QSUB', 'M111'),
    '83321-TD100NNB': ('QSUB', 'M112'), '83321-TD100PNY': ('QSUB', 'M114'),
    '83321-TD100VKE': ('QSUB', 'M116'), '83321-TD100YGN': ('QSUB', 'M117'),
    '83324-P8000': ('AQFG', 'DL002'), '83324-P8010': ('AQFG', 'DL004'),
    '83335-P8000BM1': ('Q034', '8132'), '83335-P8000DNN': ('Q034', '8133'),
    '83335-P8000JY2': ('Q034', '8175'), '83335-P8000RBQ': ('Q034', '8046'),
    '83345-P8000BM1': ('Q035', '8132'), '83345-P8000DNN': ('Q035', '8133'),
    '83345-P8000JY2': ('Q035', '8175'), '83345-P8000RBQ': ('Q035', '8046'),
    'C2311-TD000': ('QSUB', 'M087'), 'C3311-TD000': ('QSUB', 'M088'), 'C3311-TD100': ('QSUB', 'M089'),
    'C3321-TD000': ('QSUB', 'M088'), 'C3321-TD100': ('QSUB', 'M089'),
    'D0111-PI010': ('QSUB', 'N035'), 'D0121-PI010': ('QSUB', 'N035'),
    'D0131-PI020': ('QSUB', 'N037'), 'D0141-PI020': ('QSUB', 'N037'),
    'D0133-P8000': ('AQFG', 'DL007'), 'D0133-P8010': ('AQFG', 'DL009'),
    'D0143-P8000': ('AQFG', 'DL008'), 'D0143-P8010': ('AQFG', 'DL010'),
    'D3133-P8000': ('AQFG', 'DL005'), 'D3143-P8000': ('AQFG', 'DL006'),
    'M0230-P8000RBQ': ('XXXX', 'L001'), 'M0230-P8010RBQ': ('XXXX', 'L002'),
    'M0240-P8000RBQ': ('XXXX', 'L003'), 'M0240-P8010RBQ': ('XXXX', 'L004'),
    'M2310-PI010NNB': ('QSUB', 'N192'), 'M2310-PI010VKE': ('QSUB', 'N193'), 'M2310-PI010YGU': ('QSUB', 'N194'),
    'M2311-TD000NNB': ('QSUB', 'M142'), 'M2311-TD000PNY': ('QSUB', 'M144'),
    'M2311-TD000VKE': ('QSUB', 'M145'), 'M2311-TD000YGN': ('QSUB', 'M143'),
    'M2320-PI010NNB': ('QSUB', 'N198'), 'M2320-PI010VKE': ('QSUB', 'N199'), 'M2320-PI010YGU': ('QSUB', 'N200'),
    'M3310-PI020NNB': ('QSUB', 'N204'), 'M3310-PI020VKE': ('QSUB', 'N205'), 'M3310-PI020YGU': ('QSUB', 'N206'),
    'M3320-PI020NNB': ('QSUB', 'N210'), 'M3320-PI020VKE': ('QSUB', 'N211'), 'M3320-PI020YGU': ('QSUB', 'N212'),
    'M3311-TD000NNB': ('QSUB', 'M150'), 'M3311-TD000PNY': ('QSUB', 'M153'),
    'M3311-TD000VKE': ('QSUB', 'M155'), 'M3311-TD000YGN': ('QSUB', 'M151'),
    'M3311-TD100NNB': ('QSUB', 'M152'), 'M3311-TD100PNY': ('QSUB', 'M154'),
    'M3311-TD100VKE': ('QSUB', 'M156'), 'M3311-TD100YGN': ('QSUB', 'M157'),
    'M3321-TD000NNB': ('QSUB', 'M158'), 'M3321-TD000PNY': ('QSUB', 'M161'),
    'M3321-TD000VKE': ('QSUB', 'M163'), 'M3321-TD000YGN': ('QSUB', 'M159'),
    'M3321-TD100NNB': ('QSUB', 'M160'), 'M3321-TD100PNY': ('QSUB', 'M162'),
    'M3321-TD100VKE': ('QSUB', 'M164'), 'M3321-TD100YGN': ('QSUB', 'M165'),
    'M83371-P8000RBQ': ('Q034', '8106'), 'M83381-P8000RBQ': ('Q035', '8106'),
    'M83371-P8010RBQ': ('Q034', '8046'), 'M83381-P8010RBQ': ('Q035', '8046'),
}

LEVEL_COLS = ["K", "L", "M", "N", "O", "P", "Q", "R", "S"]
MCOLOR_COLS = ["Y", "Z", "AA", "AB"]
ASM_COLOR_COLS = ["F", "G", "H", "I"]
USAGE_COLS = ["BB", "BC", "BD", "BE"]       # Usage Quantity 1..4 = 변형 루트 1..4

warnings: list[str] = []


def warn(msg: str) -> None:
    if msg not in warnings:
        warnings.append(msg)


def add_once(lst: list[str], msg: str) -> None:
    if msg not in lst:
        lst.append(msg)


def s(v) -> str | None:
    if v is None:
        return None
    t = str(v).strip()
    return t or None


def is_num(v) -> bool:
    if isinstance(v, (int, float)):
        return True
    return isinstance(v, str) and re.fullmatch(r"\d+(\.\d+)?", v.strip()) is not None


def norm_color(c) -> str | None:
    c = s(c)
    if c is None or c in NO_COLOR:
        return None
    c = c.replace("_E", "")
    return COLOR_FIX.get(c, c)


# ── 1. 엑셀 → 블록 ──────────────────────────────────────────────────────────
def read_blocks(path: Path) -> list[dict]:
    wb = openpyxl.load_workbook(path, data_only=True)
    blocks: list[dict] = []
    for ws in wb.worksheets:
        section = None
        blk = None
        for r in range(1, ws.max_row + 1):
            g = lambda col: ws.cell(r, ci(col)).value  # noqa: E731
            b = s(g("B"))
            if b and b.startswith("■"):
                section = b.lstrip("■").strip()
                continue
            cname, part = s(g("C")), s(g("U"))
            if cname and cname not in ("No", "Complete ASSY (Level 0)", "#VALUE!") and part is None and s(g("E")) is None:
                blk = dict(sheet=ws.title, section=section, title=cname, roots=[], colors=[], rows=[], row=r)
                blocks.append(blk)
                continue
            if blk is None:
                continue
            root = s(g("E")) if s(g("E")) not in (None, "Part Number") else None
            if root is None:
                d = s(g("D"))
                root = d if d not in (None, "Image", "#VALUE!") else None
            if root:
                root = ROOT_FIX.get((ws.title, r), root)
                blk["roots"].append(root)
                cols = [norm_color(g(c)) for c in ASM_COLOR_COLS if g(c) is not None]
                if cols:
                    blk["colors"] = [c for c in cols if c]     # 'N/C' 만 있으면 빈 목록 = 색 없음
            if part is None or part == "Part Number":
                continue
            level = None
            for i, c in enumerate(LEVEL_COLS, 1):
                if g(c) is not None:
                    level = i
            usage = {}
            for i, uc in enumerate(USAGE_COLS):
                v = g(uc)
                if v is None:
                    continue
                if is_num(v):
                    usage[i] = float(v)
                elif s(v) == "?":
                    usage[i] = None
                # 그 밖의 문자열(변경 메모 등)은 수량이 아니다
            blk["rows"].append(dict(
                row=r, level=level, part=part, name=s(g("V")) or part, unit=s(g("W")) or "EA",
                mcolors=[norm_color(g(c)) for c in MCOLOR_COLS],
                supplier=SUPPLIER_FIX.get(s(g("AU")) or "", s(g("AU"))), notes=s(g("AY")), usage=usage))
    return blocks


# ── 2. 블록 → 품목·BOM ──────────────────────────────────────────────────────
Item = dict


def child_color(mcolors: list[str | None], n_colors: int, idx: int) -> str | None:
    vals = [c for c in mcolors[: max(n_colors, 1)] if c]
    distinct = set(vals)
    if not distinct or distinct == {"BK"}:
        return None                       # 색 무관(BK·공란) → 품번 그대로
    if len(distinct) == 1:
        return vals[0]                    # 단일 색 → 그 색으로 고정
    return mcolors[idx] if idx < len(mcolors) else None   # 조립색별 → 같은 순번의 색


def with_color(part: str, color: str | None) -> str:
    return part + color if color else part


def category_of(item_type: str, name: str, unit: str) -> str:
    if item_type == "ASSY":
        return "TRIM"
    if item_type == "SUB":
        return "SUB"
    n = name.upper()
    if unit == "SH":
        return "FABRIC"
    if "RESIN" in n:
        return "RESIN"
    if "ADHESIVE" in n or "LUBRICANT" in n:
        return "CHEM"
    if "SCREW" in n or "CLIP" in n or "FASTENER" in n:
        return "FASTENER"
    return "PART"


def build(blocks: list[dict]):
    items: dict[str, Item] = {}
    item_cars: dict[str, set] = collections.defaultdict(set)
    bom: dict[str, list] = collections.OrderedDict()     # parent ItemNo → [(child, qty, uom, note, srcrow)]
    skipped_qty: list[str] = []
    skipped_purchased: list[str] = []

    def put_item(item_no: str, name: str, item_type: str, unit: str, car: str, supplier: str | None, src: str):
        if len(item_no) > 20:
            raise SystemExit(f"ItemNo 20자 초과: {item_no} ({src})")
        item_cars[item_no].add(car)
        if item_no in items:
            cur = items[item_no]
            if cur["type"] != item_type:
                warn(f"ItemType 충돌 {item_no}: {cur['type']} vs {item_type} ({src}) — 먼저 나온 값 유지")
            if cur["name"] != name:
                warn(f"품명 상이 {item_no}: '{cur['name']}' vs '{name}' ({src}) — 먼저 나온 값 유지")
            return
        items[item_no] = dict(name=name[:MAX_NAME], type=item_type, unit=unit, supplier=supplier,
                              category=category_of(item_type, name, unit), src=src)

    for blk in blocks:
        car = SECTION_CAR[blk["section"]]
        colors = blk["colors"] or [None]
        roots = list(blk["roots"])
        # LQ2 모듈 블록: 변형 2(+CUR) 루트가 E 열에 없어 P8010 으로 유도 (DB 기존 품번과 일치)
        if blk["section"].startswith("LQ2") and roots and roots[0].startswith("M8") and len(roots) == 1:
            roots.append(roots[0].replace("P8000", "P8010"))
        n_var = len(roots)
        variants_used = {v for x in blk["rows"] for v in x["usage"]}
        for v in sorted(variants_used):
            if v >= n_var:
                warn(f"{blk['sheet']}!{blk['row']} '{blk['title']}': 변형 {v+1} 수량이 있는데 루트가 없음 — 무시(복사 잔재)")
        for vi, root in enumerate(roots):
            root_name = VARIANT_ROOT_NAME.get(root, blk["title"])
            for cidx, color in enumerate(colors):
                put_item(with_color(root, color), root_name, "ASSY", "EA", car, None, f"{blk['sheet']}!{blk['row']}")

        # 변형마다 부모 스택을 따로 둔다 — 공유 자식 행(BB·BC 둘 다)이 변형 2 의 L1 뒤에 오기 때문
        stacks = {vi: {} for vi in range(n_var)}
        prev_level = 0
        positions = collections.Counter()
        for x in blk["rows"]:
            level = x["level"]
            if level is None:            # MV1a 레진 행처럼 레벨 칸이 빈 행 → 직전 행의 자식
                level = prev_level + 1
                warn(f"{blk['sheet']}!{x['row']} {x['part']}: 레벨 없음 → L{level} 로 간주")
            prev_level = level
            src = f"{blk['sheet']}!{x['row']}"
            sup = x["supplier"]
            in_house = (sup or "").startswith(IN_HOUSE)
            # 자재 단위(G·SH)는 사내 생산품일 수 없다 — MV1a 스킨의 공급사 MIP 오기를 걸러낸다
            if in_house and x["unit"] in ("G", "SH"):
                warn(f"{src} {x['part']}: 단위 {x['unit']} 인데 공급사 {sup} → MATERIAL 로 본다(공급사 미상)")
                in_house, sup = False, None
            item_type = "SUB" if in_house else "MATERIAL"
            if sup and not in_house and sup not in VENDORS:
                raise SystemExit(f"미등록 공급사 {sup} ({src})")
            for vi in range(n_var):
                if vi not in x["usage"]:
                    continue
                qty = x["usage"][vi]
                for cidx, color in enumerate(colors):
                    child = with_color(x["part"], child_color(x["mcolors"], len(colors), cidx))
                    parent = with_color(roots[vi], color) if level == 1 else stacks[vi].get((level - 1, cidx))
                    if parent is None:
                        raise SystemExit(f"부모 없음 {src} L{level} {x['part']}")
                    put_item(child, x["name"], item_type, x["unit"], car, sup, src)
                    stacks[vi][(level, cidx)] = child
                    for k in [k for k in stacks[vi] if k[0] > level and k[1] == cidx]:
                        del stacks[vi][k]
                    if qty is None:
                        add_once(skipped_qty, f"{parent} → {child} ({src}, 수량 '?')")
                        continue
                    ptype = items[parent]["type"]
                    if ptype == "MATERIAL":
                        add_once(skipped_purchased, f"{parent} → {child} {qty:g} {x['unit']} ({src})")
                        continue
                    lines = bom.setdefault(parent, [])
                    dup = next((l for l in lines if l[0] == child), None)
                    if dup:
                        if abs(dup[1] - qty) > 1e-9:
                            warn(f"BOM 수량 상이 {parent} → {child}: {dup[1]:g} vs {qty:g} ({src}) — 먼저 나온 값 유지")
                        continue
                    positions[parent] += 10
                    lines.append((child, qty, x["unit"], x["notes"], src, positions[parent]))

    for item_no, cars in item_cars.items():
        items[item_no]["car"] = next(iter(cars)) if len(cars) == 1 else None
    return items, bom, skipped_qty, skipped_purchased


# ── 3. SQL ──────────────────────────────────────────────────────────────────
def q(v) -> str:
    if v is None:
        return "NULL"
    return "'" + str(v).replace("'", "''") + "'"


def nq(v) -> str:
    return "NULL" if v is None else "N" + q(v)


def num(v: float) -> str:
    t = f"{v:.4f}".rstrip("0").rstrip(".")
    return t if t else "0"


def bom_id(version: str, child: str) -> str:
    return "B" + hashlib.md5(f"{version}|{child}".encode()).hexdigest()[:23]


def render(items, bom, skipped_qty, skipped_purchased) -> str:
    type_count = collections.Counter(i["type"] for i in items.values())
    n_lines = sum(len(v) for v in bom.values())
    vendors_used = sorted({i["supplier"] for i in items.values() if i["supplier"] and not i["supplier"].startswith(IN_HOUSE)})
    item_vendor = [(no, VENDORS[i["supplier"]][0]) for no, i in items.items()
                   if i["supplier"] and not i["supplier"].startswith(IN_HOUSE)]

    L: list[str] = []
    w = L.append
    w("/* ------------------------------------------------------------------")
    w("   seed_md_item_bom_master_list.sql  (생성: tools/gen_md_item_bom_seed.py — 손으로 고치지 말 것)")
    w("   출처: docs/260828 BOM Master List.xlsx (시트 서베너·조지아·어번, 2026-08-28)")
    w("")
    w("   엑셀이 품목·BOM 의 유일한 정본이다. 이 스크립트는 MD_Item / MD_BomVersion / MD_Bom /")
    w("   SCM_ItemVendor 를 전부 지우고 다시 넣는다(SEMS 추출분 등 엑셀에 없는 품목은 사라진다).")
    w("   품목·BOM 을 FK 로 참조하는 행이 있으면 시작 전에 중단한다(구 FG 테이블·FK 가 남은 DB) —")
    w("   rebuild_db.sh 순서(PDA_SEED 이전)로 적용할 것. FK 없는 참조(수주·WO·LOT·재고 등)는 검사하지 않는다.")
    w("")
    w("   규칙")
    w("   · ItemNo = 품번 + 3자리 색상. 루트는 Assembly Color 마다 1행, 자식은 Material Color 가")
    w("     조립색별로 다르면 같은 순번의 색을, 단일 색이면 그 색을 붙이고 BK·공란이면 품번 그대로.")
    w("   · ItemType: Complete ASSY = ASSY / 공급사 MIP.JACKSON(자사) = SUB / LP·KD 구매 = MATERIAL.")
    w("   · RoutingType: Complete ASSY = 'A'(코어 사출 → 완제품 IMG, PP-003 이 WO 를 내는 전제) / SUB·MATERIAL = NULL.")
    w("   · ItemCategory: ASSY→TRIM, SUB→SUB, MATERIAL→FABRIC(SH)·RESIN·CHEM(접착제·윤활유)·FASTENER·PART.")
    w("   · CarType: ■ 헤더(NE1A W 는 NE1A). 여러 차종에 쓰이는 공용 자재는 NULL.")
    w("   · PGN·ALC 는 엑셀에 없어 2026-09-24 AMES_DEV 값을 되살린다(IMG 완제품 라벨용). MountPos 는 없다.")
    w("   · BOM 은 부모 품번마다 버전 1건(V-{ItemNo}-01, APPROVED, EffFrom 2026-08-28) + 직접 자식만 Level 1.")
    w("     QtyPer = Usage Quantity(변형 열), UOM = 자식 단위, ScrapPct 0, Note = 엑셀 Notes 열.")
    w("   · 구매 부품(MATERIAL) 아래의 레진은 사급이 아니므로 BOM 에 넣지 않는다(아래 목록).")
    w("   · Material Spec·Weight·Size·Grade 는 담을 컬럼이 없어 버린다(사용자 결정).")
    w("")
    w("   원본 보정")
    for (sh, r), fixed in ROOT_FIX.items():
        w(f"   · {sh}!{r} 루트 품번 → {fixed}")
    w("   · 조립색 YUG → YGU, 공급사 KD.HANHWA → KD.HANWHA, 'N/C' 는 색 없음")
    w("   · LQ2 모듈 블록의 변형 2(+CUR) 루트는 E 열에 없어 P8000 → P8010 으로 유도")
    w("   · 품번 없는 행, 변형 루트 없는 수량(조지아 NQ5a·MV1a 의 복사 잔재)은 무시")
    if skipped_qty:
        w("   BOM 제외 — 수량 미정('?'): 품목은 넣고 BOM 라인만 뺀다")
        for t in skipped_qty:
            w(f"   · {t}")
    if skipped_purchased:
        w("   BOM 제외 — 구매 부품 아래 자재(사급 아님)")
        for t in skipped_purchased:
            w(f"   · {t}")
    no_bom = [no for no, it in items.items() if it["type"] in ("ASSY", "SUB") and no not in bom]
    if no_bom:
        w("   BOM 없는 ASSY·SUB (원본에 자식 행이 없거나 품번이 비어 있음)")
        for t in sorted(no_bom):
            w(f"   · {t}")
    if warnings:
        w("   생성 경고")
        for t in warnings:
            w(f"   · {t}")
    w("")
    w(f"   결과: MD_Item {len(items)} (ASSY {type_count['ASSY']} · SUB {type_count['SUB']} · MATERIAL {type_count['MATERIAL']}),")
    w(f"         MD_BomVersion {len(bom)}, MD_Bom {n_lines}, MD_Vendor {len(vendors_used)}, SCM_ItemVendor {len(item_vendor)}")
    w("   적용: sqlcmd -S <server> -d AMES_DEV -U ames_app -P <pw> -f 65001 -I -b -i dist/seed_md_item_bom_master_list.sql")
    w("   ------------------------------------------------------------------ */")
    w("SET NOCOUNT ON;")
    w("SET XACT_ABORT ON;")
    w("")
    # FG 테이블 구성은 DB 마다 다르다(통합 마이그레이션이 DROP) — 테이블 이름을 박지 않고 FK 로 찾는다.
    w("-- ── §0 가드: 재적재 대상 밖에서 FK 로 품목·BOM 을 참조하는 행이 있으면 중단 ──────")
    w("DECLARE @guard nvarchar(max);")
    w("SELECT @guard = STRING_AGG(CAST(N'SELECT N''' + t.name + N''' WHERE EXISTS (SELECT 1 FROM '")
    w("                  + QUOTENAME(SCHEMA_NAME(t.schema_id)) + N'.' + QUOTENAME(t.name) + N')' AS nvarchar(max)), N' UNION ALL ')")
    w("FROM   sys.tables t")
    w("WHERE  t.object_id IN (SELECT fk.parent_object_id FROM sys.foreign_keys fk")
    w("                       WHERE fk.referenced_object_id IN (OBJECT_ID(N'dbo.MD_Item'), OBJECT_ID(N'dbo.MD_BomVersion'), OBJECT_ID(N'dbo.MD_Bom')))")
    w("  AND  t.object_id NOT IN (OBJECT_ID(N'dbo.MD_Item'), OBJECT_ID(N'dbo.MD_BomVersion'), OBJECT_ID(N'dbo.MD_Bom'), OBJECT_ID(N'dbo.SCM_ItemVendor'));")
    w("DECLARE @blocked TABLE (TableName sysname);")
    w("IF @guard IS NOT NULL INSERT @blocked EXEC (@guard);")
    w("IF EXISTS (SELECT 1 FROM @blocked)")
    w("BEGIN")
    w("    DECLARE @msg nvarchar(2048) = N'품목·BOM 을 FK 로 참조하는 행이 있어 전체 재적재를 할 수 없다: '")
    w("                                + (SELECT STRING_AGG(TableName, N', ') FROM @blocked);")
    w("    THROW 50001, @msg, 1;")
    w("END;")
    w("")
    w("BEGIN TRAN;")
    w("")
    w("-- ── §1 MD_Uom: SH(시트) ─────────────────────────────────────────────")
    w("INSERT INTO dbo.MD_Uom (UOMCode, UOMName, UOMCategory, BaseFlag, ConvFactor, DecimalPrec, Symbol, ActiveFlag, CreatedBy, CreatedTS)")
    w("SELECT 'SH', 'Sheet', 'QTY', 0, 1, 3, 'sh', 1, 'admin', SYSDATETIME()")
    w("WHERE  NOT EXISTS (SELECT 1 FROM dbo.MD_Uom WHERE UOMCode = 'SH');")
    w("PRINT CONCAT(N'§1 MD_Uom 추가: ', @@ROWCOUNT, N' 건');")
    w("")
    w("-- ── §2 ITEM_CATEGORY 공통코드 ───────────────────────────────────────")
    w("INSERT INTO dbo.MD_CodeItem (CodeID, GroupCode, CodeValue, CodeName, CodeNameEn, ParentCodeID, SortOrder, Attribute1, UseFlag, Description, CreatedBy, CreatedTS)")
    w("SELECT 'ITEM_CATEGORY_' + v.CodeValue, 'ITEM_CATEGORY', v.CodeValue, v.CodeName, v.CodeNameEn, NULL, v.SortOrder, v.Attribute1, 1, NULL, 'admin@ames.local', SYSDATETIME()")
    w("FROM  (VALUES")
    w(",\n".join(f"        ({q(cv)}, {nq(kn)}, {nq(en)}, {nq(attr)}, {so})" for cv, kn, en, attr, so in NEW_CATEGORIES))
    w("      ) v (CodeValue, CodeName, CodeNameEn, Attribute1, SortOrder)")
    w("WHERE  NOT EXISTS (SELECT 1 FROM dbo.MD_CodeItem c WHERE c.CodeID = 'ITEM_CATEGORY_' + v.CodeValue);")
    w("PRINT CONCAT(N'§2 ITEM_CATEGORY 추가: ', @@ROWCOUNT, N' 건');")
    w("")
    w("-- ── §3 MD_Vendor ────────────────────────────────────────────────────")
    w("INSERT INTO dbo.MD_Vendor (VendorID, VendorName, VendorType, VendorCategory, ActiveFlag, CreatedBy, CreatedTS)")
    w("SELECT v.VendorID, v.VendorName, v.VendorType, v.VendorCategory, 1, 'admin', SYSDATETIME()")
    w("FROM  (VALUES")
    w(",\n".join(f"        ({q(VENDORS[k][0])}, {nq(VENDORS[k][1])}, {q(VENDORS[k][2])}, {nq(VENDORS[k][3])})" for k in vendors_used))
    w("      ) v (VendorID, VendorName, VendorType, VendorCategory)")
    w("WHERE  NOT EXISTS (SELECT 1 FROM dbo.MD_Vendor x WHERE x.VendorID = v.VendorID);")
    w("PRINT CONCAT(N'§3 MD_Vendor 추가: ', @@ROWCOUNT, N' 건');")
    w("")
    w("-- ── §4 전체 삭제 ────────────────────────────────────────────────────")
    w("DELETE dbo.SCM_ItemVendor;  PRINT CONCAT(N'§4 SCM_ItemVendor 삭제: ', @@ROWCOUNT, N' 건');")
    w("DELETE dbo.MD_Bom;          PRINT CONCAT(N'§4 MD_Bom 삭제: ', @@ROWCOUNT, N' 건');")
    w("DELETE dbo.MD_BomVersion;   PRINT CONCAT(N'§4 MD_BomVersion 삭제: ', @@ROWCOUNT, N' 건');")
    w("DELETE dbo.MD_Item;         PRINT CONCAT(N'§4 MD_Item 삭제: ', @@ROWCOUNT, N' 건');")
    w("")
    w("-- ── §5 MD_Item ──────────────────────────────────────────────────────")
    w("INSERT INTO dbo.MD_Item (ItemNo, ItemName, ItemType, ItemCategory, CarType, DefaultUOM, RoutingType, SafetyStock, PGN, ALC, ToteFlag, ActiveFlag, CreatedBy, CreatedTS) VALUES")
    rows = []
    order = {"ASSY": 0, "SUB": 1, "MATERIAL": 2}
    for no, it in sorted(items.items(), key=lambda kv: (order[kv[1]["type"]], kv[0])):
        pgn, alc = PGN_ALC.get(no, (None, None))
        routing = "A" if it["type"] == "ASSY" else None
        rows.append(f"  ({q(no)}, {nq(it['name'])}, {q(it['type'])}, {q(it['category'])}, {q(it['car'])}, {q(it['unit'])}, {q(routing)}, 0, {q(pgn)}, {q(alc)}, 0, 1, {q(CREATED_BY)}, SYSDATETIME())")
    w(",\n".join(rows) + ";")
    w("DECLARE @n5 int = @@ROWCOUNT;")
    w(f"IF @n5 <> {len(items)} THROW 50002, N'MD_Item 행 수 불일치', 1;")
    w("PRINT CONCAT(N'§5 MD_Item 추가: ', @n5, N' 건');")
    w("")
    w("-- ── §6 MD_BomVersion (부모 품번마다 1건) ────────────────────────────")
    w("INSERT INTO dbo.MD_BomVersion (VersionID, RootItemNo, VersionNo, EffFrom, EffTo, ChangeType, ChangeReason, RequestedBy, ApprovedBy, ApprovedTS, Status, CreatedBy, CreatedTS) VALUES")
    rows = []
    for parent in bom:
        rows.append(f"  ({q('V-' + parent + '-01')}, {q(parent)}, 'V1.0', {q(EFF_FROM)}, NULL, 'INITIAL', N'BOM Master List {EFF_FROM}', {q(CREATED_BY)}, {q(CREATED_BY)}, SYSDATETIME(), 'APPROVED', {q(CREATED_BY)}, SYSDATETIME())")
    w(",\n".join(rows) + ";")
    w("DECLARE @n6 int = @@ROWCOUNT;")
    w(f"IF @n6 <> {len(bom)} THROW 50003, N'MD_BomVersion 행 수 불일치', 1;")
    w("PRINT CONCAT(N'§6 MD_BomVersion 추가: ', @n6, N' 건');")
    w("")
    w("-- ── §7 MD_Bom (직접 자식, Level 1) ──────────────────────────────────")
    w("INSERT INTO dbo.MD_Bom (BOMID, ParentItemNo, CompItemNo, BOMLevel, QtyPer, UOM, ScrapPct, VersionID, Position, Note, ActiveFlag, CreatedBy, CreatedTS) VALUES")
    rows = []
    for parent, lines in bom.items():
        ver = "V-" + parent + "-01"
        for child, qty, uom, note, src, pos in lines:
            rows.append(f"  ({q(bom_id(ver, child))}, {q(parent)}, {q(child)}, 1, {num(qty)}, {q(uom)}, 0, {q(ver)}, {pos}, {nq(note[:120] if note else None)}, 1, {q(CREATED_BY)}, SYSDATETIME())")
    w(",\n".join(rows) + ";")
    w("DECLARE @n7 int = @@ROWCOUNT;")
    w(f"IF @n7 <> {n_lines} THROW 50004, N'MD_Bom 행 수 불일치', 1;")
    w("IF EXISTS (SELECT 1 FROM dbo.MD_Bom b WHERE NOT EXISTS (SELECT 1 FROM dbo.MD_Item i WHERE i.ItemNo = b.ParentItemNo)")
    w("                                        OR NOT EXISTS (SELECT 1 FROM dbo.MD_Item i WHERE i.ItemNo = b.CompItemNo))")
    w("    THROW 50005, N'MD_Bom 에 MD_Item 에 없는 품번이 있다', 1;")
    w("PRINT CONCAT(N'§7 MD_Bom 추가: ', @n7, N' 건');")
    w("")
    w("-- ── §8 SCM_ItemVendor (구매 자재 → 공급사) ──────────────────────────")
    w("INSERT INTO dbo.SCM_ItemVendor (ItemNo, VendorID, ActiveFlag, CreatedBy, CreatedTS) VALUES")
    w(",\n".join(f"  ({q(no)}, {q(vid)}, 1, {q(CREATED_BY)}, SYSDATETIME())" for no, vid in sorted(item_vendor)) + ";")
    w("DECLARE @n8 int = @@ROWCOUNT;")
    w(f"IF @n8 <> {len(item_vendor)} THROW 50006, N'SCM_ItemVendor 행 수 불일치', 1;")
    w("PRINT CONCAT(N'§8 SCM_ItemVendor 추가: ', @n8, N' 건');")
    w("")
    w("COMMIT;")
    w("GO")
    w("")
    w("-- 확인")
    w("SELECT ItemType, COUNT(*) AS Items FROM dbo.MD_Item GROUP BY ItemType ORDER BY ItemType;")
    w("SELECT (SELECT COUNT(*) FROM dbo.MD_BomVersion) AS Versions, (SELECT COUNT(*) FROM dbo.MD_Bom) AS Lines,")
    w("       (SELECT COUNT(*) FROM dbo.MD_Item i WHERE i.ItemType IN ('ASSY','SUB')")
    w("          AND NOT EXISTS (SELECT 1 FROM dbo.MD_Bom b WHERE b.ParentItemNo = i.ItemNo)) AS MakeItemsWithoutBom;")
    w("GO")
    return "\n".join(L) + "\n"


def main() -> None:
    blocks = read_blocks(XLSX)
    items, bom, skipped_qty, skipped_purchased = build(blocks)
    sql = render(items, bom, skipped_qty, skipped_purchased)
    OUT.write_text(sql, encoding="utf-8-sig")
    print(f"blocks {len(blocks)} · items {len(items)} · versions {len(bom)} · lines {sum(len(v) for v in bom.values())}")
    print(f"skipped qty {len(skipped_qty)} · skipped purchased {len(skipped_purchased)} · warnings {len(warnings)}")
    for t in warnings:
        print("  !", t)
    print("->", OUT)


if __name__ == "__main__":
    main()
